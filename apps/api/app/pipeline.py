import io
import re
import uuid
import csv
from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
from openpyxl import load_workbook
from .schemas import Report, Section, Issue, Scores

REQUIREMENTS = {
    "resume": ["Header", "Summary", "Skills", "Experience", "Projects", "Education", "Certifications", "Achievements"],
    "research_paper": ["Title", "Abstract", "Introduction", "Methods", "Results", "Discussion", "Conclusion", "References"],
    "academic_report": ["Title", "Executive Summary", "Introduction", "Analysis", "Findings", "Conclusion", "References"],
    "project_report": ["Title", "Objectives", "Scope", "Approach", "Timeline", "Risks", "Results", "Next Steps"],
    "technical_documentation": ["Overview", "Prerequisites", "Installation", "Usage", "Configuration", "Troubleshooting", "Reference"],
    "business_proposal": ["Executive Summary", "Problem", "Proposed Solution", "Benefits", "Timeline", "Budget", "Next Steps"],
    "presentation": ["Title", "Agenda", "Context", "Key Findings", "Recommendations", "Next Steps"],
    "cover_letter": ["Header", "Opening", "Relevant Experience", "Motivation", "Closing"],
    "statement_of_purpose": ["Opening", "Academic Background", "Research Interests", "Goals", "Conclusion"],
    "general_document": ["Title", "Introduction", "Body", "Conclusion"],
}


def _norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


def _classify(text: str) -> tuple[str, float]:
    signals = {
        "resume": ["experience", "education", "skills", "work experience"],
        "research_paper": ["abstract", "methodology", "references", "results"],
        "technical_documentation": ["installation", "configuration", "api reference", "prerequisites"],
        "business_proposal": ["executive summary", "budget", "timeline", "proposed solution"],
        "project_report": ["objectives", "scope", "risks", "next steps"],
    }
    scores = {kind: sum(term in text.lower() for term in terms) for kind, terms in signals.items()}
    kind, score = max(scores.items(), key=lambda item: item[1])
    if score == 0:
        return "general_document", 0.42
    return kind, min(0.62 + score * 0.08, 0.96)


def _extract_sections(text: str) -> list[Section]:
    sections: list[Section] = []
    for index, line in enumerate(text.splitlines()):
        clean = re.sub(r"^[#\d.)\s-]+", "", line).strip()
        if 2 <= len(clean) <= 80 and (line.startswith("#") or line.isupper() or re.match(r"^\d+[.)]\s", line)):
            sections.append(Section(
                id=str(uuid.uuid4()), title=clean.title(), level=1, order=len(sections) + 1,
                start_page=max(1, text[:text.find(line)].count("\f") + 1),
                content_preview=" ".join(text.splitlines()[index + 1:index + 3])[:180],
            ))
    return sections


def extract_content(filename: str, content: bytes) -> tuple[str, int, list[str]]:
    suffix = Path(filename).suffix.lower()
    warnings: list[str] = []
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        return "\n\f\n".join(page.extract_text() or "" for page in reader.pages), len(reader.pages), warnings
    if suffix == ".docx":
        document = DocxDocument(io.BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            blocks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(blocks), 1, warnings
    if suffix == ".pptx":
        presentation = Presentation(io.BytesIO(content))
        slides = []
        for slide in presentation.slides:
            slides.append("\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text_frame")))
        return "\n\f\n".join(slides), len(presentation.slides), warnings
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheets = []
        for sheet in workbook.worksheets:
            rows = [" | ".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True)]
            sheets.append(f"{sheet.title}\n" + "\n".join(rows))
        return "\n\f\n".join(sheets), len(workbook.worksheets), warnings
    if suffix == ".csv":
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        return "\n".join(" | ".join(row) for row in rows), 1, warnings
    if suffix in {".txt", ".md"}:
        return content.decode("utf-8-sig"), 1, warnings
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        raise ValueError("Image uploads require an OCR worker. Configure an OCR provider before uploading image-only documents.")
    raise ValueError(f"Unsupported document type: {suffix or 'unknown'}")


def analyze_document(filename: str, content: bytes) -> Report:
    raw, pages, warnings = extract_content(filename, content)
    if not raw.strip():
        raise ValueError("The document contains no extractable text. An OCR-enabled ingestion worker is required for scanned documents.")
    document_type, confidence = _classify(raw)
    sections = _extract_sections(raw)
    actual = [_norm(section.title) for section in sections]
    required = REQUIREMENTS[document_type]
    missing = [name for name in required if _norm(name) not in actual]
    ordering: list[Issue] = []
    positions = {key: actual.index(_norm(key)) for key in required if _norm(key) in actual}
    present = [key for key in required if _norm(key) in positions]
    if present != sorted(present, key=lambda key: required.index(key)):
        ordering.append(Issue(
            id=str(uuid.uuid4()), category="ordering", severity="high",
            explanation="Detected sections do not follow the expected order for this document type.",
            evidence=["Actual: " + " → ".join(section.title for section in sections)],
            recommendation="Move the sections into the recommended outline shown below.",
            confidence=0.91,
        ))
    content_issues: list[Issue] = []
    if document_type == "resume" and re.search(r"\b(experience|projects)\b", raw, re.I) and not re.search(r"\b\d+%|\$\d+|\b\d+\s+(users|clients|projects)\b", raw, re.I):
        content_issues.append(Issue(
            id=str(uuid.uuid4()), category="content_quality", severity="medium",
            explanation="Experience evidence lacks measurable outcomes.",
            evidence=["No percentages, currency values, or quantified outcomes were detected in the extracted text."],
            recommendation="Rewrite key bullets with an action, scope, and measurable result.",
            confidence=0.87, affected_section="Experience",
        ))
    completeness = round((len(required) - len(missing)) / len(required) * 100)
    structure = max(0, 100 - len(ordering) * 18)
    quality = max(0, 100 - len(content_issues) * 12)
    compliance = round((completeness + structure) / 2)
    recommendations = [f"Add the missing {name} section." for name in missing] + [issue.recommendation for issue in ordering + content_issues]
    return Report(
        document_id=str(uuid.uuid4()), filename=filename, document_type=document_type, confidence=confidence,
        title=sections[0].title if sections else filename.rsplit(".", 1)[0], pages=pages,
        sections=sections, missing_sections=missing, ordering_issues=ordering, content_issues=content_issues,
        recommendations=recommendations, restructured_outline=required,
        scores=Scores(overall=round((completeness + structure + quality + compliance) / 4),
                      completeness=completeness, structure=structure, content_quality=quality, compliance=compliance),
        processing={"stages": ["Uploading", "Parsing", "Understanding", "Classifying", "Checking requirements", "Finding gaps", "Evaluating quality", "Generating recommendations", "Complete"], "warnings": warnings},
    )


analyze_pdf = analyze_document
